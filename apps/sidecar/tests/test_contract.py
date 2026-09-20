import hashlib
import io
import json
import math
import os
import subprocess
import sys
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
import requests
import tldextract
import yaml
from docx import Document
from pydantic import TypeAdapter, ValidationError

from redactio_sidecar.engine import Engine
from redactio_sidecar.ipc import EngineError, dispatch_request, run_loop
from redactio_sidecar.schemas import (
    Decisions,
    Detection,
    ProcessingConfig,
    Request,
    Response,
    ReviewRequest,
)

MODEL_NAME = "OpenMed-PII-German-BiomedBERT-Large-340M-v1"

PAIR_ID = "11111111-1111-4111-8111-111111111111"
REVISION = "22222222-2222-4222-8222-222222222222"
RULE_ID = "33333333-3333-4333-8333-333333333333"
TIMESTAMP = "2026-09-19T10:00:00+02:00"
REQUEST_ADAPTER = TypeAdapter(Request)
RESPONSE_ADAPTER = TypeAdapter(Response)


def engine_info() -> dict[str, Any]:
    return {
        "engine_version": "synthetic-engine",
        "model_name": MODEL_NAME,
        "model_version": "synthetic-model",
        "recognizers": ["synthetic-recognizer"],
        "extraction_version": "synthetic-extractor",
    }


def document_meta() -> dict[str, Any]:
    return {
        "sync_pair_id": PAIR_ID,
        "doc_id": "doc-0001",
        "source_hash_sha256": "0" * 64,
        "processing_revision": REVISION,
        "redacted_at": TIMESTAMP,
    }


def detection() -> dict[str, Any]:
    return {
        "id": "detection-1",
        "start": 0,
        "end": 4,
        "entity_type": "PERSON",
        "confidence": 0.9,
        "recognizer": "synthetic-recognizer",
        "origin": "automatic",
    }


def process_result_payload() -> dict[str, Any]:
    return {
        **document_meta(),
        "markdown": "---\nschema_version: 1\n---\n<PERSON_1>",
        "body": "<PERSON_1>",
        "original_text": "Anna",
        "detections": [detection()],
        "redactions": [
            {
                "start_offset": 0,
                "end_offset": 10,
                "entity_type": "PERSON",
                "placeholder": "<PERSON_1>",
                "confidence": 0.9,
                "recognizer": "synthetic-recognizer",
                "origin": "automatic",
            }
        ],
        "warnings": [],
        "body_was_empty": False,
        "review_status": "pending",
        "engine": engine_info(),
    }


def synthetic_messages() -> list[dict[str, Any]]:
    config = {
        "model": MODEL_NAME,
        "enabled_entities": [
            "PERSON",
            "LOCATION",
            "EMAIL_ADDRESS",
            "PHONE_NUMBER",
            "IBAN_CODE",
            "IP_ADDRESS",
            "URL",
            "DATE_TIME",
        ],
        "custom_rules": [
            {
                "id": RULE_ID,
                "entity_type": "CUSTOM",
                "enabled": True,
                "kind": "regex",
                "pattern": "Synthetic_[0-9]+",
            },
            {
                "id": "44444444-4444-4444-8444-444444444444",
                "entity_type": "CUSTOM",
                "enabled": True,
                "kind": "words",
                "words": ["Synthetic"],
            },
        ],
        "include_positions": True,
    }
    review_payload = {
        **document_meta(),
        "source_path": "synthetic.docx",
        "detections": [detection()],
        "decisions": {"dismissed_ids": [], "manual": []},
        "review_status": "pending",
        "reviewed_at": None,
        "acknowledged_warnings": [],
    }
    return [
        {"id": "ping", "type": "ping", "payload": {}},
        {
            "id": "configure",
            "type": "configure",
            "payload": {
                "sync_pair_id": PAIR_ID,
                "processing_revision": REVISION,
                "config": config,
            },
        },
        {
            "id": "process",
            "type": "process_document",
            "payload": {**document_meta(), "source_path": "synthetic.docx"},
        },
        {
            "id": "preview",
            "type": "preview_rules",
            "payload": {
                "sync_pair_id": PAIR_ID,
                "processing_revision": REVISION,
                "text": "Synthetic_1",
            },
        },
        {"id": "render", "type": "render_review", "payload": review_payload},
        {
            "id": "ping",
            "type": "ping_result",
            "payload": {"protocol_version": 1},
        },
        {
            "id": "configure",
            "type": "configure_result",
            "payload": {
                "sync_pair_id": PAIR_ID,
                "processing_revision": REVISION,
                "engine": engine_info(),
            },
        },
        {
            "id": "process",
            "type": "process_document_result",
            "payload": process_result_payload(),
        },
        {
            "id": "preview",
            "type": "preview_rules_result",
            "payload": {
                "sync_pair_id": PAIR_ID,
                "processing_revision": REVISION,
                "detections": [detection()],
            },
        },
        {
            "id": "render",
            "type": "render_review_result",
            "payload": process_result_payload(),
        },
        {
            "id": "error",
            "type": "error",
            "payload": {"code": "synthetic_error", "retryable": False},
        },
    ]


def test_processing_defaults_and_empty_decisions_match_shared_contract():
    config = ProcessingConfig()
    assert config.model == MODEL_NAME
    assert config.enabled_entities == [
        "PERSON",
        "LOCATION",
        "EMAIL_ADDRESS",
        "PHONE_NUMBER",
        "IBAN_CODE",
        "IP_ADDRESS",
        "URL",
        "DATE_TIME",
    ]
    assert config.model_entities is None
    assert config.custom_rules == []
    assert config.include_positions is True
    assert Decisions().model_dump() == {"dismissed_ids": [], "manual": []}


@pytest.mark.parametrize(
    "mutation",
    [
        lambda value: value.update(sync_pair_id="not-a-uuid"),
        lambda value: value.update(processing_revision="not-a-uuid"),
        lambda value: value.update(doc_id="doc-1"),
        lambda value: value.update(source_hash_sha256="A" * 64),
        lambda value: value.update(redacted_at="2026-09-19T10:00:00"),
        lambda value: value.update(extra="private"),
    ],
)
def test_document_metadata_rejects_invalid_fields(mutation):
    payload = document_meta()
    mutation(payload)
    message = {
        "id": "invalid",
        "type": "process_document",
        "payload": {**payload, "source_path": "synthetic.docx"},
    }
    with pytest.raises(ValidationError):
        REQUEST_ADAPTER.validate_python(message)


@pytest.mark.parametrize(
    "mutation",
    [
        lambda value: value.update(start=-1),
        lambda value: value.update(end=0),
        lambda value: value.update(entity_type="not_valid"),
        lambda value: value.update(confidence=-0.1),
        lambda value: value.update(confidence=1.1),
        lambda value: value.update(confidence=math.nan),
        lambda value: value.update(extra="private"),
    ],
)
def test_detection_rejects_invalid_spans_confidence_and_fields(mutation):
    value = detection()
    mutation(value)
    with pytest.raises(ValidationError):
        Detection.model_validate(value)


@pytest.mark.parametrize("confidence", [True, "0.5"])
def test_detection_rejects_non_numeric_json_confidence(confidence):
    value = detection()
    value["confidence"] = confidence
    with pytest.raises(ValidationError):
        Detection.model_validate(value)


@pytest.mark.parametrize("field", ["enabled", "include_positions"])
def test_processing_config_rejects_coerced_json_booleans(field):
    value: dict[str, Any] = {"include_positions": True}
    if field == "enabled":
        value["custom_rules"] = [
            {
                "id": RULE_ID,
                "entity_type": "CUSTOM",
                "enabled": "false",
                "kind": "words",
                "words": ["Synthetic"],
            }
        ]
    else:
        value[field] = "false"
    with pytest.raises(ValidationError):
        ProcessingConfig.model_validate(value)


@pytest.mark.parametrize("timestamp", [0, "2026-09-19 10:00:00+02:00"])
def test_document_metadata_requires_rfc3339_timestamp_string(timestamp):
    payload = document_meta()
    payload["redacted_at"] = timestamp
    message = {
        "id": "invalid-time",
        "type": "process_document",
        "payload": {**payload, "source_path": "synthetic.docx"},
    }
    with pytest.raises(ValidationError):
        REQUEST_ADAPTER.validate_python(message)


def test_document_metadata_accepts_lowercase_rfc3339_separators():
    payload = document_meta()
    payload["redacted_at"] = "2026-09-19t10:00:00z"
    request = REQUEST_ADAPTER.validate_python(
        {
            "id": "lowercase-time",
            "type": "process_document",
            "payload": {**payload, "source_path": "synthetic.docx"},
        }
    )
    assert request.payload.redacted_at == datetime(2026, 9, 19, 10, 0, tzinfo=UTC)


def test_internal_model_construction_accepts_aware_datetime():
    payload = document_meta()
    payload["redacted_at"] = datetime(2026, 9, 19, 8, 0, tzinfo=UTC)
    REQUEST_ADAPTER.validate_python(
        {
            "id": "internal-time",
            "type": "process_document",
            "payload": {**payload, "source_path": "synthetic.docx"},
        }
    )


@pytest.mark.parametrize(
    "rule",
    [
        {
            "id": RULE_ID,
            "entity_type": "CUSTOM",
            "enabled": True,
            "kind": "regex",
            "pattern": "(",
        },
        {
            "id": RULE_ID,
            "entity_type": "CUSTOM",
            "enabled": True,
            "kind": "words",
            "words": [""],
        },
    ],
)
def test_custom_rules_reject_invalid_regex_and_empty_words(rule):
    with pytest.raises(ValidationError):
        ProcessingConfig(custom_rules=[rule])


def test_processing_config_rejects_duplicate_rule_ids():
    rule = {
        "id": RULE_ID,
        "entity_type": "CUSTOM",
        "enabled": True,
        "kind": "words",
        "words": ["Synthetic"],
    }
    with pytest.raises(ValidationError):
        ProcessingConfig(custom_rules=[rule, deepcopy(rule)])


def test_all_synthetic_contract_messages_validate_and_reject_extra_fields():
    messages = synthetic_messages()
    for message in messages[:5]:
        REQUEST_ADAPTER.validate_python(message)
    for message in messages[5:]:
        RESPONSE_ADAPTER.validate_python(message)

    invalid = deepcopy(messages[0])
    invalid["payload"]["extra"] = "private"
    with pytest.raises(ValidationError):
        REQUEST_ADAPTER.validate_python(invalid)


@pytest.mark.parametrize("message_index,required_field", [(1, "config"), (4, "decisions")])
def test_request_payloads_require_all_shared_contract_fields(message_index, required_field):
    invalid = deepcopy(synthetic_messages()[message_index])
    del invalid["payload"][required_field]
    with pytest.raises(ValidationError):
        REQUEST_ADAPTER.validate_python(invalid)


def test_emit_cli_outputs_only_valid_jsonl_contract_objects():
    completed = subprocess.run(
        [sys.executable, str(Path(__file__).resolve()), "--emit"],
        check=True,
        capture_output=True,
        text=True,
    )
    emitted = [json.loads(line) for line in completed.stdout.splitlines()]
    assert emitted == synthetic_messages()


def write_document(path: Path, text: str, *, header: str | None = None) -> str:
    document = Document()
    if text:
        document.add_paragraph(text)
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    document.save(path)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def model_root() -> Path:
    root = Path(os.environ.get("REDACTIO_MODEL_DIR", Path(__file__).parents[1] / "models"))
    assert (root / "manifest.json").is_file()
    return root


def processing_config(*, include_positions: bool = False) -> dict[str, Any]:
    return {
        "model": MODEL_NAME,
        "enabled_entities": [],
        "custom_rules": [
            {
                "id": RULE_ID,
                "entity_type": "CUSTOM",
                "enabled": True,
                "kind": "regex",
                "pattern": "Synthetic_[0-9]+",
            }
        ],
        "include_positions": include_positions,
    }


def test_real_jsonl_process_and_review_are_private_and_repeatable(tmp_path, capsys):
    path = tmp_path / "CANARY_PATIENT_FILE.docx"
    source_hash = write_document(path, "Synthetic_1\n---\nraw: Käthe")
    meta = {
        "sync_pair_id": PAIR_ID,
        "doc_id": "doc-0001",
        "source_hash_sha256": source_hash,
        "processing_revision": REVISION,
        "redacted_at": TIMESTAMP,
        "source_path": str(path),
    }
    manual = {
        "id": "manual-1",
        "start": 0,
        "end": 11,
        "entity_type": "CUSTOM",
        "confidence": None,
        "recognizer": "manual",
        "origin": "manual",
    }
    review = {
        **meta,
        "detections": [],
        "decisions": {"dismissed_ids": [], "manual": [manual]},
        "review_status": "approved",
        "reviewed_at": TIMESTAMP,
        "acknowledged_warnings": [],
    }
    messages = [
        {"id": "ping", "type": "ping", "payload": {}},
        {
            "id": "configure",
            "type": "configure",
            "payload": {
                "sync_pair_id": PAIR_ID,
                "processing_revision": REVISION,
                "config": processing_config(),
            },
        },
        {"id": "process", "type": "process_document", "payload": meta},
        {"id": "render-1", "type": "render_review", "payload": review},
        {"id": "render-2", "type": "render_review", "payload": review},
        {
            "id": "changed",
            "type": "render_review",
            "payload": {
                **review,
                "source_hash_sha256": "0" * 64,
                "decisions": {"dismissed_ids": ["unknown"], "manual": []},
            },
        },
    ]
    stdin = io.BytesIO(b"".join(json.dumps(message).encode() + b"\n" for message in messages))
    stdout = io.BytesIO()
    engine = Engine(model_root())

    run_loop(stdin, stdout, lambda request: dispatch_request(request, engine))

    replies = [json.loads(line) for line in stdout.getvalue().splitlines()]
    assert [reply["type"] for reply in replies] == [
        "ping_result",
        "configure_result",
        "process_document_result",
        "render_review_result",
        "render_review_result",
        "error",
    ]
    processed = replies[2]["payload"]
    assert processed["original_text"] == "Synthetic_1\n---\nraw: Käthe"
    assert processed["body"] == "<CUSTOM_1>\n---\nraw: Käthe\n"
    assert processed["markdown"].partition("\n---\n")[2] == processed["body"]
    assert processed["redactions"]
    assert "redactions:" not in processed["markdown"]
    rendered = replies[3]["payload"]
    assert rendered["review_status"] == "approved"
    assert rendered["body"] == "<CUSTOM_1>\n---\nraw: Käthe\n"
    assert rendered["markdown"].partition("\n---\n")[2] == rendered["body"]
    assert rendered["markdown"] == replies[4]["payload"]["markdown"]
    frontmatter = yaml.safe_load(rendered["markdown"].split("---", 2)[1])
    assert frontmatter["reviewed_at"] == TIMESTAMP
    assert replies[5]["payload"] == {"code": "source_changed", "retryable": False}
    assert b"CANARY_PATIENT_FILE" not in stdout.getvalue()
    assert "CANARY_PATIENT_FILE" not in capsys.readouterr().err


def test_review_uses_stored_spans_and_enforces_warning_approval(tmp_path):
    engine = Engine(model_root())
    engine.configure(PAIR_ID, REVISION, ProcessingConfig.model_validate(processing_config()))
    path = tmp_path / "warning.docx"
    source_hash = write_document(path, "Synthetic_1", header="private header")
    request = ReviewRequest.model_validate(
        {
            **document_meta(),
            "source_hash_sha256": source_hash,
            "source_path": str(path),
            "detections": [],
            "decisions": {"dismissed_ids": [], "manual": []},
            "review_status": "approved",
            "reviewed_at": TIMESTAMP,
            "acknowledged_warnings": [],
        }
    )

    with pytest.raises(EngineError, match="approval_not_allowed"):
        engine.render_review(request)

    approved = engine.render_review(
        request.model_copy(update={"acknowledged_warnings": ["headers_footers"]})
    )
    assert approved.body == "Synthetic_1\n"
    assert approved.markdown.partition("\n---\n")[2] == approved.body
    assert approved.redactions == []
    assert approved.review_status == "approved"
    assert approved.warnings == ["headers_footers"]

    with pytest.raises(EngineError, match="approval_not_allowed"):
        engine.render_review(
            request.model_copy(
                update={
                    "acknowledged_warnings": ["headers_footers"],
                    "reviewed_at": None,
                }
            )
        )

    invalid_detection = Detection(
        id="stored-manual",
        start=0,
        end=11,
        entity_type="CUSTOM",
        confidence=None,
        recognizer="manual",
        origin="manual",
    )
    with pytest.raises(EngineError, match="invalid_review"):
        engine.render_review(
            request.model_copy(
                update={
                    "detections": [invalid_detection],
                    "review_status": "rejected",
                }
            )
        )

    generated = engine.analyze(PAIR_ID, REVISION, "Synthetic_1")
    with pytest.raises(EngineError, match="invalid_review"):
        engine.render_review(
            request.model_copy(
                update={
                    "detections": [generated[0].model_copy(update={"id": "tampered"})],
                    "review_status": "rejected",
                }
            )
        )


def test_review_allows_disabled_native_labels_for_manual_detections(tmp_path):
    engine = Engine(model_root())
    engine.configure(
        PAIR_ID,
        REVISION,
        ProcessingConfig(model=MODEL_NAME, enabled_entities=[], model_entities=[]),
    )
    path = tmp_path / "native-manual.docx"
    source_hash = write_document(path, "42")
    request = ReviewRequest.model_validate(
        {
            **document_meta(),
            "source_hash_sha256": source_hash,
            "source_path": str(path),
            "detections": [],
            "decisions": {
                "dismissed_ids": [],
                "manual": [
                    {
                        "id": "manual-age",
                        "start": 0,
                        "end": 2,
                        "entity_type": "AGE",
                        "confidence": None,
                        "recognizer": "manual",
                        "origin": "manual",
                    }
                ],
            },
            "review_status": "pending",
            "reviewed_at": None,
            "acknowledged_warnings": [],
        }
    )

    assert engine.render_review(request).body == "<AGE_1>\n"
    with pytest.raises(EngineError, match="invalid_review"):
        engine.render_review(
            request.model_copy(
                update={
                    "decisions": request.decisions.model_copy(
                        update={
                            "manual": [
                                request.decisions.manual[0].model_copy(
                                    update={"entity_type": "NOT_A_MODEL_LABEL"}
                                )
                            ]
                        }
                    )
                }
            )
        )


def test_empty_processing_needs_rework_and_cannot_be_approved(tmp_path):
    engine = Engine(model_root())
    engine.configure(PAIR_ID, REVISION, ProcessingConfig.model_validate(processing_config()))
    path = tmp_path / "empty.docx"
    source_hash = write_document(path, "")
    payload = {
        **document_meta(),
        "source_hash_sha256": source_hash,
        "source_path": str(path),
    }

    processed = engine.process_document(
        REQUEST_ADAPTER.validate_python(
            {"id": "empty", "type": "process_document", "payload": payload}
        ).payload
    )
    assert processed.body == "\n"
    assert processed.markdown.partition("\n---\n")[2] == processed.body
    assert processed.body_was_empty is True
    assert processed.review_status == "needs-rework"

    review = ReviewRequest.model_validate(
        {
            **payload,
            "detections": [],
            "decisions": {"dismissed_ids": [], "manual": []},
            "review_status": "approved",
            "reviewed_at": TIMESTAMP,
            "acknowledged_warnings": ["empty_document"],
        }
    )
    with pytest.raises(EngineError, match="approval_not_allowed"):
        engine.render_review(review)


def test_cli_takes_explicit_model_root_and_keeps_ping_available_without_models(tmp_path):
    messages = [
        {"id": "ping", "type": "ping", "payload": {}},
        {
            "id": "configure",
            "type": "configure",
            "payload": {
                "sync_pair_id": PAIR_ID,
                "processing_revision": REVISION,
                "config": processing_config(),
            },
        },
    ]
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "redactio_sidecar",
            "--model-dir",
            str(tmp_path / "missing-models"),
        ],
        input="".join(json.dumps(message) + "\n" for message in messages),
        capture_output=True,
        text=True,
        check=True,
    )

    replies = [json.loads(line) for line in completed.stdout.splitlines()]
    assert replies[0]["type"] == "ping_result"
    assert replies[1]["payload"] == {
        "code": "invalid_model_manifest",
        "retryable": False,
    }
    assert completed.stderr == ""


def test_first_email_analysis_uses_bundled_suffix_data_without_network(
    tmp_path, monkeypatch, capsys, caplog
):
    attempts = []

    def reject_network(_session, request, **_kwargs):
        attempts.append(request.url)
        raise requests.ConnectionError("network disabled")

    monkeypatch.setattr(
        tldextract,
        "extract",
        tldextract.TLDExtract(cache_dir=str(tmp_path / "cold-cache")),
    )
    monkeypatch.setattr(requests.sessions.Session, "send", reject_network)
    engine = Engine(model_root())
    engine.configure(
        PAIR_ID,
        REVISION,
        ProcessingConfig(enabled_entities=["EMAIL_ADDRESS"]),
    )

    detections = engine.analyze(PAIR_ID, REVISION, "Kontakt: kontakt@example.org")

    assert {detection.entity_type for detection in detections} == {"EMAIL_ADDRESS"}
    assert attempts == []
    assert capsys.readouterr().err == ""
    assert not [record for record in caplog.records if record.name.startswith("tldextract")]


if __name__ == "__main__" and sys.argv[1:] == ["--emit"]:
    for synthetic_message in synthetic_messages():
        print(json.dumps(synthetic_message, separators=(",", ":")))
