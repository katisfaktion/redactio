"""Generate synthetic packaging input; never overwrite an existing document."""

import argparse
import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document


def canary(index: int):
    document = Document()
    for text in (
        f"Synthetischer Fall {index}.",
        "Kontakt: anna.beispiel@example.invalid",
        "Max Mustermann besucht uns. Der Termin ist in Berlin.",
        "Kontakt: max@example.com",
        "Telefon: +49 30 12345678",
        "IBAN: DE89370400440532013000",
        "Server: 192.168.1.1",
        "Webseite: https://example.de/path",
        "Termin: 19.09.2026",
    ):
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Kategorie"
    table.cell(0, 1).text = "Synthetischer Inhalt"
    return document


def save(document, path: Path) -> None:
    # python-docx otherwise gives ZIP members the wall-clock save time.
    document.core_properties.created = datetime(2026, 9, 19, tzinfo=UTC)
    document.core_properties.modified = document.core_properties.created
    buffer = BytesIO()
    document.save(buffer)
    with ZipFile(buffer) as original, ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name in sorted(original.namelist()):
            info = ZipInfo(name, date_time=(2026, 9, 19, 0, 0, 0))
            info.create_system = 3  # Keep identical archive headers on Windows and POSIX.
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, original.read(name))


def prepare(output: Path) -> None:
    if output.is_symlink() or (output.exists() and any(output.iterdir())):
        raise ValueError("The output directory must be empty and dedicated")


def manifest(output: Path, entries: list[dict]) -> None:
    for entry in entries:
        entry["sha256"] = hashlib.sha256((output / entry["path"]).read_bytes()).hexdigest()
    (output / "expectations.json").write_text(
        json.dumps({"schema_version": 2, "documents": len(entries), "files": entries}, indent=2)
        + "\n",
        encoding="utf-8",
    )


def generate(output: Path, count: int, edge_output: Path | None = None) -> None:
    if not 1 <= count <= 10000:
        raise ValueError("Count must be between 1 and 10000")
    prepare(output)
    if edge_output is not None:
        prepare(edge_output)
        if (
            output.resolve() == edge_output.resolve()
            or output.resolve() in edge_output.resolve().parents
            or edge_output.resolve() in output.resolve().parents
        ):
            raise ValueError("Edge corpus must be a separate directory")
    output.mkdir(parents=True, exist_ok=True)
    entries = []
    for index in range(1, count + 1):
        name = f"case-{index:04}.docx"
        save(canary(index), output / name)
        entries.append({"path": name, "profile": "canary", "warnings": []})
    manifest(output, entries)
    if edge_output is None:
        return
    edge_output.mkdir(parents=True, exist_ok=True)
    entries = []
    for profile in ("warnings", "corrupt", "empty", "repeated", "unicode", "tamper"):
        name = f"{profile}.docx"
        document = Document() if profile == "empty" else canary(1)
        warnings = []
        if profile == "warnings":
            document.sections[0].header.paragraphs[0].text = "Synthetische Kopfzeile"
            warnings = ["headers_footers"]
        elif profile == "empty":
            warnings = ["empty_document"]
        elif profile == "repeated":
            document.add_paragraph("Max Mustermann. Max Mustermann.")
        elif profile == "unicode":
            document.add_paragraph("Ä Ö Ü ß 🙂 e\u0301 — synthetisch")
        if profile == "corrupt":
            (edge_output / name).write_bytes(b"PK\x03\x04synthetic corrupt ZIP")
        else:
            save(document, edge_output / name)
        entry = {"path": name, "profile": profile, "warnings": warnings}
        if profile == "repeated":
            entry["configuration"] = "literal_person_rule"
        if profile == "corrupt":
            entry["error"] = "invalid_docx"
        if profile == "tamper":
            # Applied only AFTER real host output exists, by the acceptance operator.
            entry["after_processing"] = "append_output_then_scan"
            entry["expected_state"] = "conflict"
        entries.append(entry)
    manifest(edge_output, entries)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--count", default=400, type=int)
    parser.add_argument("--edge-output", type=Path)
    args = parser.parse_args()
    generate(args.output, args.count, args.edge_output)
